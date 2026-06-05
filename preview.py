import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont
prs = Presentation("/home/user/Soutenance/Soutenance_Aphelie_Ilyesse_Kebaili.pptx")
SCALE = 96/914400
W = int(prs.slide_width*SCALE); H = int(prs.slide_height*SCALE)
want = [int(a) for a in sys.argv[1:]] if len(sys.argv)>1 else None
def font(sz,bold=False):
    p="/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    try: return ImageFont.truetype(p,max(8,int(sz)))
    except: return ImageFont.load_default()
def rgb(c):
    try: return (c[0],c[1],c[2])
    except: return None
def fc(sh):
    try:
        if sh.fill.type==1: return rgb(sh.fill.fore_color.rgb)
    except: pass
    try:
        from pptx.oxml.ns import qn
        g=sh.fill._xPr.findall('.//'+qn('a:srgbClr'))
        if g: v=g[0].get('val'); return (int(v[0:2],16),int(v[2:4],16),int(v[4:6],16))
    except: pass
    return None
def lc(sh):
    try:
        if sh.line.fill.type==1: return rgb(sh.line.color.rgb)
    except: pass
    return None
def wrap(d,t,f,mw):
    out=[];cur=""
    for w in t.split():
        s=(cur+" "+w).strip()
        if d.textlength(s,font=f)<=mw or not cur: cur=s
        else: out.append(cur);cur=w
    if cur:out.append(cur)
    return out
def rc(r):
    try: return rgb(r.font.color.rgb)
    except: return (30,30,30)
def draw_sh(d,sh,ox=0,oy=0):
    x=int(sh.left*SCALE)+ox;y=int(sh.top*SCALE)+oy
    w=int(sh.width*SCALE);h=int(sh.height*SCALE);x1=x+w;y1=y+h
    if sh.shape_type==MSO_SHAPE_TYPE.TABLE:
        tb=sh.table
        rows=list(tb.rows);cols=list(tb.columns)
        ys=[y]
        for r in rows: ys.append(ys[-1]+int(r.height*SCALE))
        xs=[x]
        for c in cols: xs.append(xs[-1]+int(c.width*SCALE))
        for ri in range(len(rows)):
            for ci in range(len(cols)):
                cell=tb.cell(ri,ci)
                cx0,cy0,cx1,cy1=xs[ci],ys[ri],xs[ci+1],ys[ri+1]
                cf=None
                try:
                    if cell.fill.type==1: cf=rgb(cell.fill.fore_color.rgb)
                except: pass
                d.rectangle([cx0,cy0,cx1,cy1],fill=cf,outline=(215,224,232))
                cyt=cy0+3
                for para in cell.text_frame.paragraphs:
                    runs=para.runs
                    if not runs: continue
                    sz=runs[0].font.size.pt if runs[0].font.size else 11
                    f=font(sz,any(r.font.bold for r in runs))
                    full="".join(r.text for r in runs)
                    for ln in wrap(d,full,f,max(10,cx1-cx0-8)):
                        d.text((cx0+4,cyt),ln,font=f,fill=rc(runs[0]) or (30,30,30));cyt+=int(sz*1.25)+1
        return
    f0=fc(sh);l0=lc(sh)
    try: stp=str(sh.auto_shape_type)
    except: stp=None
    if stp:
        if 'OVAL' in stp: d.ellipse([x,y,x1,y1],fill=f0,outline=l0)
        elif 'HEXAGON' in stp:
            inset=w*0.25
            d.polygon([(x+inset,y),(x1-inset,y),(x1,(y+y1)/2),(x1-inset,y1),(x+inset,y1),(x,(y+y1)/2)],fill=f0,outline=l0)
        elif f0 or l0: d.rounded_rectangle([x,y,x1,y1],radius=6 if 'ROUNDED' in stp else 0,fill=f0,outline=l0)
    if sh.has_text_frame:
        cyt=y+4
        for para in sh.text_frame.paragraphs:
            runs=para.runs
            if not runs: continue
            sz=runs[0].font.size.pt if runs[0].font.size else 14
            f=font(sz,any(r.font.bold for r in runs))
            full="".join(r.text for r in runs)
            for ln in wrap(d,full,f,max(10,w-8)):
                tw=d.textlength(ln,font=f);al=str(para.alignment)
                tx=x+(w-tw)/2 if 'CENTER' in al else (x1-tw-4 if 'RIGHT' in al else x+4)
                d.text((tx,cyt),ln,font=f,fill=rc(runs[0]) or (30,30,30));cyt+=int(sz*1.25)+2
            cyt+=int(para.space_after.pt if para.space_after else 4)
for idx,slide in enumerate(prs.slides,1):
    if want and idx not in want: continue
    img=Image.new("RGB",(W,H),(255,255,255));d=ImageDraw.Draw(img)
    for sh in slide.shapes: draw_sh(d,sh)
    img.save(f"/home/user/Soutenance/prev_{idx:02d}.png")
print("ok",want)
